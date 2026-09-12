"""
Deliverables: a multi-sheet workbook and a set of publication figures.

The figures are the point. A results chapter that says "C++ was faster" is an
assertion; a latency distribution with confidence intervals, a stage-attribution
waterfall and a scaling curve are evidence. Everything here is written at 300
dpi with a consistent, print-safe style so the files can go straight into a
report without being redrawn.

Both openpyxl and matplotlib are optional. Missing either degrades that half of
the output and says so, rather than losing the run.
"""

from __future__ import annotations

import csv
import statistics
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from aoi.analysis import StudyResult, Summary, summarise
from aoi.pipeline import STAGES

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    _XLSX = True
except ImportError:
    _XLSX = False

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    _MPL = True
except ImportError:
    _MPL = False


def xlsx_available() -> bool:
    return _XLSX


def figures_available() -> bool:
    return _MPL


# =============================================================================
# Figure style
# =============================================================================

# Print-safe and distinguishable in greyscale, which matters because a thesis
# gets photocopied. Path A warm, Path B cool, ablation muted.
COLOURS = {
    "python": "#D1495B",
    "python-lean": "#E8A33D",
    "cpp": "#2E6F95",
    "sim": "#8D99AE",
}
STAGE_COLOURS = [
    "#2E6F95", "#4E96B8", "#7FB8CE", "#E8A33D",
    "#D1495B", "#8D5A97", "#5B8C5A", "#B0B7C3",
]


def _style():
    plt.rcParams.update({
        "figure.dpi": 120, "savefig.dpi": 300,
        "font.family": "DejaVu Sans", "font.size": 10,
        "axes.titlesize": 12, "axes.titleweight": "bold",
        "axes.labelsize": 10, "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "grid.alpha": 0.22, "grid.linewidth": 0.6,
        "axes.axisbelow": True, "legend.frameon": False,
        "figure.facecolor": "white", "savefig.bbox": "tight",
    })


def _colour(engine: str) -> str:
    return COLOURS.get(engine, "#6C757D")


def _save(fig, out_dir: Path, name: str, written: List[Path]):
    path = out_dir / f"{name}.png"
    fig.savefig(path)
    # PDF as well: vector output is what a printed report actually wants.
    fig.savefig(out_dir / f"{name}.pdf")
    plt.close(fig)
    written.append(path)


# =============================================================================
# Figures
# =============================================================================

def figure_latency_distribution(study: StudyResult, out_dir: Path,
                                written: List[Path]) -> None:
    """Per-frame latency distribution with the mean and its 95% CI marked.

    A box plot alone hides multimodality - a path that alternates between two
    latency regimes looks identical to one with a wide unimodal spread. The
    strip of individual frames behind the box makes that visible.
    """
    engines = [e for e in study.engines() if study.latencies(e)]
    if not engines:
        return
    fig, ax = plt.subplots(figsize=(1.9 * len(engines) + 3.2, 4.6))

    import numpy as np
    rng = np.random.default_rng(3)
    ymax = max(max(study.latencies(e)) for e in engines)
    # Headroom reserved up front so the per-engine annotations sit in clear space
    # instead of colliding with the title when one path dominates the scale.
    ax.set_ylim(0, ymax * 1.30)
    label_y = ymax * 1.08

    for i, e in enumerate(engines):
        vals = study.latencies(e)
        s = summarise(vals)
        ax.scatter(np.full(len(vals), i) + rng.normal(0, 0.055, len(vals)), vals,
                   s=11, alpha=0.28, color=_colour(e), linewidths=0, zorder=2)
        bp = ax.boxplot([vals], positions=[i], widths=0.42, showfliers=False,
                        patch_artist=True, zorder=3)
        bp["boxes"][0].set(facecolor="white", edgecolor=_colour(e), linewidth=1.6,
                           alpha=0.9)
        for part in ("whiskers", "caps", "medians"):
            for a in bp[part]:
                a.set(color=_colour(e), linewidth=1.6)
        ax.errorbar(i, s.mean, yerr=[[s.mean - s.ci95[0]], [s.ci95[1] - s.mean]],
                    fmt="D", ms=6, color=_colour(e), capsize=5, capthick=1.8,
                    elinewidth=1.8, zorder=4, markeredgecolor="white")
        ax.annotate(f"{s.mean:.1f} ms\n{1000 / s.mean if s.mean else 0:.1f} FPS",
                    (i, label_y), ha="center", va="bottom", fontsize=9,
                    color=_colour(e), weight="bold")

    ax.set_xticks(range(len(engines)))
    ax.set_xticklabels([study.label(e).replace(" - ", "\n") for e in engines])
    ax.set_ylabel("Frame latency (ms)")
    ax.set_title("Per-frame latency distribution")
    n = sum(len(study.latencies(e)) for e in engines)
    ax.text(0.99, 0.02,
            f"{study.config.trials} trials x {study.config.trial_frames} frames "
            f"({n} measured)   diamond = mean, bar = 95% CI",
            transform=ax.transAxes, ha="right", va="bottom", fontsize=8,
            color="#5A6472")
    _save(fig, out_dir, "fig1_latency_distribution", written)


def figure_stage_breakdown(study: StudyResult, out_dir: Path,
                           written: List[Path]) -> None:
    """Where each path spends its frame, stage by stage.

    Stacked absolute milliseconds rather than percentages: the interesting
    comparison is how much time a stage costs, not what fraction of a total that
    itself differs between paths.
    """
    engines = [e for e in study.engines() if study.latencies(e)]
    if not engines:
        return
    fig, ax = plt.subplots(figsize=(9.2, 0.7 * len(engines) + 2.1))
    ax.set_ylim(len(engines) - 0.5, -0.5)

    labels = [s.replace("_ms", "") for s in STAGES]
    left = [0.0] * len(engines)
    for si, stage in enumerate(STAGES):
        vals = [study.stage_means(e)[stage] for e in engines]
        if all(v < 0.005 for v in vals):
            continue
        ax.barh(range(len(engines)), vals, left=left, height=0.46,
                color=STAGE_COLOURS[si % len(STAGE_COLOURS)],
                edgecolor="white", linewidth=0.8, label=labels[si])
        for i, v in enumerate(vals):
            if v > max(sum(study.stage_means(e).values()) for e in engines) * 0.055:
                ax.text(left[i] + v / 2, i, f"{v:.1f}", ha="center", va="center",
                        fontsize=8.5, color="white", weight="bold")
            left[i] += v

    ax.set_yticks(range(len(engines)))
    ax.set_yticklabels([study.label(e) for e in engines])
    ax.set_xlabel("Mean time per frame (ms)")
    ax.set_title("Frame time decomposition by pipeline stage")
    ax.legend(ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.16), fontsize=9)
    _save(fig, out_dir, "fig2_stage_breakdown", written)


def figure_stage_attribution(study: StudyResult, out_dir: Path,
                             written: List[Path], a: str = "python",
                             b: str = "cpp") -> None:
    """Waterfall: which stages the A-over-B gap is actually made of.

    This is the figure that turns the headline ratio into a claim about
    mechanism. A bar dominated by `marshal` supports the data-handoff argument;
    one dominated by `inference` says the difference is runtime, not language,
    and the write-up has to say so.
    """
    if not (study.latencies(a) and study.latencies(b)):
        return
    rows = [r for r in study.stage_attribution(a, b) if abs(r["delta_ms"]) > 0.005]
    if not rows:
        return

    fig, ax = plt.subplots(figsize=(8.4, 4.4))
    names = [r["stage"] for r in rows]
    deltas = [r["delta_ms"] for r in rows]
    colours = ["#D1495B" if d > 0 else "#2E6F95" for d in deltas]
    bars = ax.bar(names, deltas, color=colours, edgecolor="white", linewidth=0.8)

    for bar, r in zip(bars, rows):
        d = r["delta_ms"]
        ax.annotate(f"{d:+.1f} ms\n{r['share'] * 100:.0f}%",
                    (bar.get_x() + bar.get_width() / 2, d),
                    textcoords="offset points",
                    xytext=(0, 6 if d >= 0 else -22),
                    ha="center", fontsize=8.5, weight="bold",
                    color="#2B3038")

    ax.axhline(0, color="#2B3038", linewidth=1.0)
    ax.set_ylabel(f"{study.label(a)} minus {study.label(b)}  (ms)")
    ax.set_title("Attribution of the latency gap by stage")
    total = sum(deltas)
    ax.text(0.99, 0.96, f"total gap {total:+.1f} ms/frame",
            transform=ax.transAxes, ha="right", va="top", fontsize=9,
            color="#5A6472")
    pad = max(abs(min(deltas)), abs(max(deltas))) * 0.28
    ax.set_ylim(min(0, min(deltas)) - pad, max(0, max(deltas)) + pad)
    _save(fig, out_dir, "fig3_stage_attribution", written)


def figure_resources(study: StudyResult, out_dir: Path,
                     written: List[Path]) -> None:
    """Throughput next to the hardware it costs.

    Latency alone would let a path that pins every core look like a free win.
    """
    engines = [e for e in study.engines() if study.latencies(e)]
    if not engines:
        return
    panels = [("fps", "Effective FPS", "", lambda e: 1000.0 / study.summary(e).mean
               if study.summary(e).mean else 0.0),
              ("cpu_percent", "CPU load", "%",
               lambda e: _mean(study.resource(e, "cpu_percent"))),
              ("gpu_percent", "GPU load", "%",
               lambda e: _mean(study.resource(e, "gpu_percent"))),
              ("rss_mb", "Process RSS", "MB",
               lambda e: _mean(study.resource(e, "rss_mb")))]

    fig, axes = plt.subplots(1, len(panels), figsize=(3.1 * len(panels), 3.9))
    for ax, (_key, title, unit, getter) in zip(axes, panels):
        vals = [getter(e) for e in engines]
        bars = ax.bar(range(len(engines)), vals,
                      color=[_colour(e) for e in engines],
                      edgecolor="white", linewidth=0.8, width=0.62)
        for bar, v in zip(bars, vals):
            ax.annotate(f"{v:.0f}" if v >= 10 else f"{v:.2f}",
                        (bar.get_x() + bar.get_width() / 2, v),
                        textcoords="offset points", xytext=(0, 4),
                        ha="center", fontsize=9, weight="bold")
        ax.set_xticks(range(len(engines)))
        ax.set_xticklabels([e for e in engines], rotation=18, ha="right", fontsize=9)
        ax.set_title(f"{title}" + (f" ({unit})" if unit else ""))
        ax.set_ylim(0, max(vals) * 1.25 if max(vals) else 1)
    fig.suptitle("Throughput and hardware utilisation", y=1.02,
                 fontsize=12, fontweight="bold")
    _save(fig, out_dir, "fig4_resources", written)


def figure_timeline(study: StudyResult, out_dir: Path,
                    written: List[Path]) -> None:
    """Latency against frame index: the tail behaviour a mean cannot show.

    Periodic spikes on the interpreted path are the visible signature of
    garbage collection and allocator behaviour; a flat compiled trace next to it
    is a stronger argument than any summary statistic.
    """
    engines = [e for e in study.engines() if len(study.latencies(e)) > 3]
    if not engines:
        return
    fig, ax = plt.subplots(figsize=(9.4, 4.0))
    for e in engines:
        vals = study.latencies(e)
        ax.plot(range(len(vals)), vals, lw=1.3, color=_colour(e),
                label=study.label(e), alpha=0.92)
        s = summarise(vals)
        ax.axhline(s.mean, color=_colour(e), lw=0.9, ls="--", alpha=0.55)
    ax.set_xlabel("Measured frame index (trials concatenated)")
    ax.set_ylabel("Latency (ms)")
    ax.set_title("Per-frame latency trace")
    ax.set_ylim(bottom=0)
    ax.legend(ncol=len(engines), loc="upper center", bbox_to_anchor=(0.5, -0.17))
    _save(fig, out_dir, "fig5_timeline", written)


def figure_sweep_tiles(rows: Sequence[Dict[str, object]], out_dir: Path,
                       written: List[Path]) -> None:
    """Latency and per-tile cost against the SAHI grid.

    The left panel answers "how much does slicing cost"; the right panel is the
    one that separates fixed overhead from per-call overhead. A path whose
    ms-per-tile falls as the grid grows is amortising a fixed cost; one that
    stays flat is paying per dispatch.
    """
    rows = [r for r in rows if "error" not in r]
    if not rows:
        return
    engines = sorted({r["engine"] for r in rows})
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(11.4, 4.3))

    for e in engines:
        pts = sorted([r for r in rows if r["engine"] == e], key=lambda r: r["tiles"])
        if not pts:
            continue
        x = [p["tiles"] for p in pts]
        y = [p["mean_ms"] for p in pts]
        lo = [p["ci_lo"] for p in pts]
        hi = [p["ci_hi"] for p in pts]
        ax1.plot(x, y, "o-", color=_colour(e), lw=1.8, ms=6, label=e)
        ax1.fill_between(x, lo, hi, color=_colour(e), alpha=0.16, linewidth=0)
        ax2.plot(x, [p["ms_per_tile"] for p in pts], "s--", color=_colour(e),
                 lw=1.6, ms=5, label=e)

    for ax, ylab, title in ((ax1, "Frame latency (ms)", "Latency vs SAHI grid"),
                            (ax2, "ms per tile", "Per-tile cost (amortisation)")):
        ax.set_xlabel("Tiles per frame")
        ax.set_ylabel(ylab)
        ax.set_title(title)
        ax.set_ylim(bottom=0)
        ax.legend()
    _save(fig, out_dir, "fig6_sweep_tiles", written)


def figure_sweep_resolution(rows: Sequence[Dict[str, object]], out_dir: Path,
                            written: List[Path]) -> None:
    """Scaling with input resolution, in ms and in ms per megapixel."""
    rows = [r for r in rows if "error" not in r]
    if not rows:
        return
    engines = sorted({r["engine"] for r in rows})
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(11.4, 4.3))

    for e in engines:
        pts = sorted([r for r in rows if r["engine"] == e],
                     key=lambda r: r["megapixels"])
        if not pts:
            continue
        x = [p["megapixels"] for p in pts]
        ax1.plot(x, [p["mean_ms"] for p in pts], "o-", color=_colour(e),
                 lw=1.8, ms=6, label=e)
        ax1.fill_between(x, [p["ci_lo"] for p in pts], [p["ci_hi"] for p in pts],
                         color=_colour(e), alpha=0.16, linewidth=0)
        ax2.plot(x, [p["ms_per_megapixel"] for p in pts], "s--", color=_colour(e),
                 lw=1.6, ms=5, label=e)

    labels = sorted({(r["megapixels"], f"{r['width']}x{r['height']}") for r in rows})
    for ax, ylab, title in ((ax1, "Frame latency (ms)", "Latency vs input resolution"),
                            (ax2, "ms per megapixel", "Normalised pixel cost")):
        ax.set_xlabel("Input (megapixels)")
        ax.set_ylabel(ylab)
        ax.set_title(title)
        ax.set_ylim(bottom=0)
        ax.set_xticks([m for m, _ in labels])
        ax.set_xticklabels([f"{n}\n{m:.1f} MP" for m, n in labels], fontsize=8)
        ax.legend()
    _save(fig, out_dir, "fig7_sweep_resolution", written)


def figure_accuracy(study: StudyResult, out_dir: Path,
                    written: List[Path]) -> None:
    """Detection quality per path, against the synthetic ground truth.

    Without this the latency comparison is unsafe: a path that finds fewer
    defects will always look faster. Equal precision/recall is what licenses the
    speed claim.
    """
    engines = [e for e in study.engines() if study.accuracy(e)]
    if not engines:
        return
    keys = ("precision", "recall", "f1", "ap50", "mean_iou")
    fig, ax = plt.subplots(figsize=(1.5 * len(engines) * len(keys) / 2 + 4.0, 4.2))

    width = 0.8 / len(engines)
    import numpy as np
    base = np.arange(len(keys))
    for i, e in enumerate(engines):
        acc = study.accuracy(e)
        vals = [acc.get(k, 0.0) for k in keys]
        bars = ax.bar(base + i * width - 0.4 + width / 2, vals, width * 0.92,
                      color=_colour(e), edgecolor="white", linewidth=0.8,
                      label=study.label(e))
        for bar, v in zip(bars, vals):
            ax.annotate(f"{v:.2f}", (bar.get_x() + bar.get_width() / 2, v),
                        textcoords="offset points", xytext=(0, 3),
                        ha="center", fontsize=8)
    ax.set_xticks(base)
    ax.set_xticklabels([k.upper() if k == "ap50" else k.replace("_", " ").title()
                        for k in keys])
    ax.set_ylim(0, 1.12)
    ax.set_ylabel("Score")
    ax.set_title("Detection quality vs planted ground truth (IoU 0.5)")
    ax.legend(ncol=len(engines))
    _save(fig, out_dir, "fig8_accuracy", written)


def _mean(vals: Sequence[float]) -> float:
    return statistics.fmean(vals) if vals else 0.0


def write_figures(study: StudyResult, out_dir: str | Path) -> List[Path]:
    """Render every figure the available data supports."""
    if not _MPL:
        return []
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    _style()
    written: List[Path] = []

    figure_latency_distribution(study, out, written)
    figure_stage_breakdown(study, out, written)
    figure_stage_attribution(study, out, written)
    figure_resources(study, out, written)
    figure_timeline(study, out, written)
    figure_accuracy(study, out, written)
    if study.sweeps.get("tiles"):
        figure_sweep_tiles(study.sweeps["tiles"], out, written)
    if study.sweeps.get("resolution"):
        figure_sweep_resolution(study.sweeps["resolution"], out, written)
    return written


# =============================================================================
# Workbook
# =============================================================================

_HDR_FILL = "1E293B"
_HDR_FONT = "FFFFFF"
_ACCENT = "E8EEF9"
FONT = "Calibri"


def _header(ws, row: int, ncols: int, first_col: int = 1):
    thin = Side(style="thin", color="94A3B8")
    for c in range(first_col, first_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color=_HDR_FONT)
        cell.fill = PatternFill("solid", fgColor=_HDR_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)


def _title(ws, text: str, subtitle: str = ""):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=FONT, italic=True, size=9, color="7B8794")


def _autosize(ws, min_w=11, max_w=52):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None),
                      default=0)
        ws.column_dimensions[letter].width = max(min_w, min(max_w, longest + 3))


def _table(ws, row: int, headers: Sequence[str],
           rows: Sequence[Sequence], formats: Optional[Sequence[str]] = None) -> int:
    for i, h in enumerate(headers, start=1):
        ws.cell(row=row, column=i, value=h)
    _header(ws, row, len(headers))
    row += 1
    for r in rows:
        for i, v in enumerate(r, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            cell.font = Font(name=FONT)
            if formats and i - 1 < len(formats) and formats[i - 1]:
                cell.number_format = formats[i - 1]
        row += 1
    return row + 1


def write_workbook(study: StudyResult, path: str | Path) -> Path:
    """Multi-sheet workbook: environment, statistics, attribution, sweeps, raw."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not _XLSX:
        return _write_csv_fallback(study, path.with_suffix(".csv"))

    wb = Workbook()
    cfg = study.config

    # ---------------- 1. Method & environment ----------------
    ws = wb.active
    ws.title = "Method"
    _title(ws, "AOI Comparative Study - Method and Environment",
           f"Generated {datetime.now():%Y-%m-%d %H:%M:%S} | "
           f"config fingerprint {cfg.fingerprint()}")

    row = 4
    ws.cell(row=row, column=1, value="Measurement protocol").font = Font(
        name=FONT, bold=True, size=11)
    row += 1
    protocol = [
        ("Trials per engine", cfg.trials),
        ("Frames per trial", cfg.trial_frames),
        ("Warmup passes (discarded)", cfg.warmup_runs),
        ("Settling frames discarded per trial", cfg.discard_frames),
        ("Measured frames per engine", cfg.trials * max(0, cfg.trial_frames - cfg.discard_frames)),
        ("Input", f"{cfg.frame_width}x{cfg.frame_height} synthetic board, "
                  f"{cfg.synth_defects} planted defects, seed {cfg.synth_seed}"),
        ("SAHI grid", f"{cfg.tile_width}x{cfg.tile_height}, "
                      f"{cfg.overlap_ratio:.0%} requested overlap"),
        ("Model input", f"{cfg.input_size}x{cfg.input_size}"),
        ("Thresholds", f"conf {cfg.conf_threshold}, NMS IoU {cfg.nms_threshold}"),
        ("CUDA / FP16", f"{cfg.use_cuda} / {cfg.use_fp16}"),
        ("Study duration", f"{study.duration_s:.1f} s"),
    ]
    row = _table(ws, row, ["Parameter", "Value"], protocol)

    ws.cell(row=row, column=1, value="Environment").font = Font(
        name=FONT, bold=True, size=11)
    row += 1
    row = _table(ws, row, ["Item", "Value"],
                 [(k, str(v)) for k, v in study.environment.items()])

    if study.notes:
        ws.cell(row=row, column=1, value="Notes and warnings").font = Font(
            name=FONT, bold=True, size=11)
        row += 1
        for n in study.notes:
            ws.cell(row=row, column=1, value=n).font = Font(name=FONT, size=9)
            row += 1
    _autosize(ws)

    # ---------------- 2. Results ----------------
    ws2 = wb.create_sheet("Results")
    _title(ws2, "Latency statistics per test path",
           "Interval estimates. A mean without its CI is not a result.")
    engines = [e for e in study.engines() if study.latencies(e)]
    rows = []
    for e in engines:
        s = study.summary(e)
        rows.append([study.label(e), study.backend(e), s.n, s.mean, s.sd,
                     s.ci95[0], s.ci95[1], s.median, s.p95, s.p99, s.cv,
                     1000.0 / s.mean if s.mean else 0.0,
                     study.between_trial_cv(e)])
    fmt = [None, None, "0", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00",
           "0.00", "0.0%", "0.00", "0.0%"]
    row = _table(ws2, 4,
                 ["Path", "Backend", "n", "Mean (ms)", "SD (ms)", "CI95 lo",
                  "CI95 hi", "Median", "p95", "p99", "CV", "FPS",
                  "Between-trial CV"],
                 rows, fmt)

    cmp_ = study.headline()
    if cmp_:
        ws2.cell(row=row, column=1, value="Hypothesis test").font = Font(
            name=FONT, bold=True, size=11)
        row += 1
        row = _table(ws2, row, ["Statistic", "Value"], [
            ("Comparison", f"{cmp_.a_label} vs {cmp_.b_label}"),
            ("Test", "Welch's unequal-variance t-test, two-sided"),
            ("t", round(cmp_.t_statistic, 4)),
            ("df", round(cmp_.df, 2)),
            ("p-value", cmp_.p_value),
            ("Significant at 0.05", "yes" if cmp_.significant else "no"),
            ("Cliff's delta", round(cmp_.cliffs_delta, 4)),
            ("Effect size", cmp_.effect_label),
            ("Speedup (mean ratio)", round(cmp_.speedup, 4)),
            ("Speedup CI95", f"[{cmp_.speedup_ci95[0]:.3f}, "
                             f"{cmp_.speedup_ci95[1]:.3f}] (bootstrap, 4000 resamples)"),
            ("Verdict", cmp_.verdict()),
        ])
        ws2.cell(row=row - 2, column=2).fill = PatternFill("solid", fgColor=_ACCENT)

    par = study.parity_summary()
    if par:
        ws2.cell(row=row, column=1, value="Detection parity (Path A vs Path B)").font = \
            Font(name=FONT, bold=True, size=11)
        row += 1
        row = _table(ws2, row, ["Metric", "Value"],
                     [(k, round(v, 4)) for k, v in par.items()])
        ws2.cell(row=row, column=1, value=(
            "Agreement is the Jaccard index of IoU-matched boxes at 0.5. A speed "
            "claim is only valid while this stays high: a path that detects less "
            "is trivially faster.")).font = Font(name=FONT, italic=True, size=9)
        row += 2
    _autosize(ws2)

    # ---------------- 3. Stage attribution ----------------
    ws3 = wb.create_sheet("Stage attribution")
    _title(ws3, "Where the latency gap comes from",
           "Per-frame means. 'share' is each stage's fraction of the total gap.")
    if len(engines) >= 2:
        a, b = ("python", "cpp") if {"python", "cpp"} <= set(engines) else engines[:2]
        attr = study.stage_attribution(a, b)
        row = _table(ws3, 4,
                     ["Stage", f"{study.label(a)} (ms)", f"{study.label(b)} (ms)",
                      "Delta (ms)", "Share of gap", "Ratio A/B"],
                     [[r["stage"], r["a_ms"], r["b_ms"], r["delta_ms"],
                       r["share"], r["ratio"]] for r in attr],
                     [None, "0.000", "0.000", "0.000", "0.0%", "0.00"])
        ws3.cell(row=row, column=1, value=(
            "marshal = cross-language pixel handoff; dispatch = residual framework "
            "overhead, computed as the frame total minus every named stage, so it "
            "cannot be silently absorbed elsewhere.")).font = Font(
            name=FONT, italic=True, size=9)

    row = (row if len(engines) >= 2 else 4) + 2
    ws3.cell(row=row, column=1, value="Mean stage times, all paths").font = Font(
        name=FONT, bold=True, size=11)
    row += 1
    _table(ws3, row, ["Stage"] + [study.label(e) for e in engines],
           [[s.replace("_ms", "")] + [study.stage_means(e)[s] for e in engines]
            for s in STAGES],
           [None] + ["0.000"] * len(engines))
    _autosize(ws3)

    # ---------------- 4. Accuracy ----------------
    acc_engines = [e for e in engines if study.accuracy(e)]
    if acc_engines:
        ws4 = wb.create_sheet("Accuracy")
        _title(ws4, "Detection quality vs planted ground truth",
               "Synthetic board: defect coordinates are known exactly, so "
               "precision and recall are computable with no annotated dataset.")
        keys = ["tp", "fp", "fn", "precision", "recall", "f1", "mean_iou", "ap50"]
        _table(ws4, 4, ["Path"] + [k.upper() if k in ("tp", "fp", "fn", "ap50")
                                   else k.title() for k in keys],
               [[study.label(e)] + [study.accuracy(e).get(k, 0.0) for k in keys]
                for e in acc_engines],
               [None, "0.0", "0.0", "0.0", "0.000", "0.000", "0.000",
                "0.000", "0.000"])
        _autosize(ws4)

    # ---------------- 5. Sweeps ----------------
    for key, title in (("tiles", "Sweep - SAHI grid"),
                       ("resolution", "Sweep - input resolution")):
        rows_ = study.sweeps.get(key)
        if not rows_:
            continue
        wsx = wb.create_sheet(title[:31])
        _title(wsx, title, "One operating point cannot show whether the gap is "
                           "fixed overhead or scales with the work.")
        cols = [c for c in rows_[0] if c != "error"]
        _table(wsx, 4, cols, [[r.get(c) for c in cols] for r in rows_])
        _autosize(wsx)

    # ---------------- 6. Raw ----------------
    ws6 = wb.create_sheet("Raw frames")
    headers = ["engine", "trial", "frame", "total_ms", "fps"] + list(STAGES)
    ws6.append(headers)
    _header(ws6, 1, len(headers))
    for t in study.trials:
        for i, lat in enumerate(t.latencies):
            ws6.append([t.engine, t.trial, i, lat,
                        1000.0 / lat if lat else 0.0]
                       + [t.stages.get(s, [0.0] * len(t.latencies))[i]
                          if i < len(t.stages.get(s, [])) else 0.0 for s in STAGES])
    ws6.freeze_panes = "A2"
    _autosize(ws6)

    wb.save(path)
    return path


def _write_csv_fallback(study: StudyResult, path: Path) -> Path:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["AOI comparative study", datetime.now().isoformat(timespec="seconds")])
        w.writerow([])
        w.writerow(["engine", "n", "mean_ms", "ci95_lo", "ci95_hi", "p95_ms",
                    "cv", "fps"])
        for e in study.engines():
            s = study.summary(e)
            w.writerow([e, s.n, f"{s.mean:.3f}", f"{s.ci95[0]:.3f}",
                        f"{s.ci95[1]:.3f}", f"{s.p95:.3f}", f"{s.cv:.4f}",
                        f"{1000 / s.mean if s.mean else 0:.2f}"])
        cmp_ = study.headline()
        if cmp_:
            w.writerow([])
            w.writerow(["verdict", cmp_.verdict()])
        w.writerow([])
        w.writerow(["engine", "trial", "frame", "total_ms"] + list(STAGES))
        for t in study.trials:
            for i, lat in enumerate(t.latencies):
                w.writerow([t.engine, t.trial, i, f"{lat:.3f}"]
                           + [f"{t.stages[s][i]:.3f}" if i < len(t.stages.get(s, []))
                              else "" for s in STAGES])
    return path


# =============================================================================
# Operator inspection report (the shop-floor artefact, not the study)
# =============================================================================

def write_inspection_report(out_dir: str | Path, latest: Dict,
                            history: Sequence[Dict]) -> Path:
    """Traceable record of a shift: yield, defect mix, per-board history."""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if not _XLSX:
        path = out / f"inspection_{stamp}.csv"
        with open(path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["metric", "value"])
            for k, v in latest.items():
                if not isinstance(v, dict):
                    w.writerow([k, v])
            w.writerow([])
            if history:
                cols = list(history[0])
                w.writerow(cols)
                for h in history:
                    w.writerow([h.get(c) for c in cols])
        return path

    path = out / f"inspection_{stamp}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Inspection"
    _title(ws, "AOI Inspection Report", f"Generated {datetime.now():%Y-%m-%d %H:%M:%S}")

    total = len(history)
    passed = sum(1 for h in history if not h.get("total_defects"))
    row = _table(ws, 4, ["Metric", "Value"], [
        ("Boards inspected", total),
        ("Boards passed", passed),
        ("Boards failed", total - passed),
    ])
    # Yield as a live formula so the sheet still recalculates if a reviewer
    # edits a count during sign-off.
    ws.cell(row=row - 2, column=1, value="Inspection yield")
    c = ws.cell(row=row - 2, column=2, value="=IFERROR(B6/B5,0)")
    c.number_format = "0.0%"
    c.fill = PatternFill("solid", fgColor=_ACCENT)

    counts: Dict[str, int] = {}
    for h in history:
        for k, v in (h.get("defect_counts") or {}).items():
            counts[k] = counts.get(k, 0) + v
    if counts:
        row = _table(ws, row + 1, ["Defect type", "Count"],
                     sorted(counts.items(), key=lambda kv: -kv[1]))
    _autosize(ws)

    if history:
        ws2 = wb.create_sheet("Boards")
        cols = ["timestamp", "engine", "total_defects", "top_defect_type",
                "latency_ms", "fps", "tiles", "cpu_percent", "gpu_percent", "rss_mb"]
        ws2.append(cols)
        _header(ws2, 1, len(cols))
        for h in history:
            ws2.append([h.get(c) for c in cols])
        ws2.freeze_panes = "A2"
        _autosize(ws2)

    wb.save(path)
    return path
